import sqlite3
from openpyxl import load_workbook
import os

# ===== BANCO =====
conn = sqlite3.connect("database\\database\\alunosv3.db")
cursor = conn.cursor()

# # # # # ===== TABELA (UMA VEZ) =====
# # # # cursor.execute(
# # # #     """
# # # # CREATE TABLE alunosImportados (
# # # #     id INTEGER PRIMARY KEY AUTOINCREMENT,
# # # #     ra TEXT,
# # # #     rm TEXT,
# # # #     responsavel_matricula TEXT,
# # # #     nome_completo TEXT,
# # # #     data_nascimento TEXT,
# # # #     rg_aluno TEXT,
# # # #     cpf_aluno TEXT,
# # # #     uf_aluno TEXT,
# # # #     expedicao_aluno TEXT,
# # # #     rne_aluno TEXT
# # # # );
# # # # """
# # # # )

# ===== PASTA =====
pasta = "primeirosAlunos"

# ===== INSERT =====
sql_insert = """
INSERT INTO alunosImportados (
    ra, rm, responsavel_matricula, nome_completo, data_nascimento,
    rg_aluno, cpf_aluno, uf_aluno, expedicao_aluno, rne_aluno,
    cartao_sus, certidao_nascimento, municipio_nascimento, distrito,
    expedicao_rg_aluno, nacionalidade, sexo, raca_cor,
    tipo_sanguineo, alergico, diabetico, observacoes,
    nome_mae, profissao_mae, email_mae, cpf_mae, rg_mae,
    uf_rg_mae, expedicao_rg_mae, estado_civil_mae,
    nome_pai, profissao_pai,
    telefone1, telefone2, telefone3, telefone4,
    bolsa_familia, nis, plano,
    escola_anterior, ano, cidade, estado,
    logradouro, numero, bairro, cidade_endereco, cep,
    nivel
) VALUES (
    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
    ?, ?, ?, ?, ?, ?, ?, ?,
    ?, ?, ?, ?,
    ?, ?, ?, ?, ?, ?, ?, ?,
    ?, ?,
    ?, ?, ?, ?,
    ?, ?, ?,
    ?, ?, ?, ?,
    ?, ?, ?, ?, ?,
    ?
);
"""
contador = 0

for arquivo in os.listdir(pasta):
    if arquivo.endswith(".xlsx"):
        caminho = os.path.join(pasta, arquivo)
        wb = load_workbook(caminho, data_only=True)
        ws = wb.active

        valores = (
            ws["O10"].value,  # RA
            ws["D10"].value,  # RM
            ws["B29"].value,  # RESPONSÁVEL PELA MATRÍCULA
            ws["B16"].value,  # NOME COMPLETO
            ws["B18"].value,  # DATA DE NASCIMENTO
            ws["B48"].value,  # RG ALUNO
            ws["B50"].value,  # CPF ALUNO
            ws["T48"].value,  # UF ALUNO
            ws["W48"].value,  # EXPEDIÇÃO ALUNO
            ws["AD48"].value,  # RNE ALUNO
            ws["AD50"].value,  # CARTÃO SUS
            ws["AB46"].value,  # CERTIDÃO DE NASCIMENTO
            ws["AD46"].value,  # MUNICÍPIO DE NASCIMENTO
            ws["AN46"].value,  # DISTRITO
            ws["W48"].value,  # EXPEDIÇÃO RG ALUNO
            ws["AE20"].value,  # NACIONALIDADE
            ws["L18"].value,  # SEXO
            ws["B20"].value,  # RAÇA/COR
            # ws["B22"].value, #PESO
            # ws["B24"].value, #ALTURA
            # ws["B26"].value, #UNIFORME
            # ws["B28"].value, #CALÇADO
            ws["AT55"].value,  # TIPO SANGUÍNEO
            ws["B55"].value,  # ALÉRGICO
            ws["B59"].value,  # DIABÉTICO
            ws["B62"].value,  # OBSERVAÇÕES
            ws["B25"].value,  # NOME MÃE
            ws["AF25"].value,  # PROFISSÃO MÃE
            ws["V29"].value,  # EMAIL MÃE
            ws["AF31"].value,  # CPF MÃE
            ws["B31"].value,  # RG MÃE
            ws["T31"].value,  # UF RG MÃE
            ws["W31"].value,  # EXPEDIÇÃO RG MÃE
            ws["AX31"].value,  # ESTADO CIVIL MÃE
            ws["B27"].value,  # NOME PAI
            ws["AF27"].value,  # PROFISSÃO PAI
            ws["B33"].value,  # TELEFONE 1
            ws["O33"].value,  # TELEFONE 2
            ws["AC33"].value,  # TELEFONE 3
            ws["AR33"].value,  # TELEFONE 4
            ws["B52"].value,  # BOLSA FAMÍLIA
            ws["N52"].value,  # NIS
            ws["AD59"].value,  # PLANO
            ws["B69"].value,  # ESCOLA ANTERIOR
            ws["AF69"].value,  # ANO
            ws["AJ69"].value,  # CIDADE
            ws["BA69"].value,  # ESTADO
            ws["B37"].value,  # LOGRADOURO
            ws["AE37"].value,  # NÚMERO
            ws["B39"].value,  # BAIRRO
            ws["AJ39"].value,  # CIDADE ENDEREÇO
            ws["R39"].value,  # CEP
            ws["K72"].value,  # NÍVEL
        )

        cursor.execute(sql_insert, valores)
        contador += 1

conn.commit()
conn.close()

print(f"{contador} alunos importados com sucesso.")
