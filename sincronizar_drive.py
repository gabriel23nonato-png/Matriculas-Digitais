#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sincronizar_drive.py
────────────────────
Registra no sistema os documentos que JÁ estão no Google Drive, sem baixar nem
reenviar nada. O sistema só considera um documento "entregue" quando existe uma
linha na tabela `documento`; este script cria essas linhas a partir das pastas

    MATRICULAS_ESCOLA / ANO / SALA / ALUNOS / NOME DO ALUNO / SUBPASTA / arquivo

Uso (na pasta do projeto, ao lado do app.py e do drive_service.py):

    python sincronizar_drive.py             # 1) SIMULA: não altera o banco;
                                            #    gera sincronizacao_drive.xlsx
    (abra a planilha, revise/complete as linhas marcadas e salve)
    python sincronizar_drive.py --aplicar   # 2) grava no banco o que está na planilha

Pode ser executado quantas vezes quiser: arquivos já registrados são ignorados.
Antes de gravar, é feito um backup automático do banco.
"""

import argparse
import difflib
import os
import re
import shutil
import sqlite3
import sys
import unicodedata
from datetime import datetime


def app_dir() -> str:
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = app_dir()
sys.path.insert(0, BASE_DIR)

from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(BASE_DIR, ".env"))

try:
    import drive_service as drive  # noqa: E402
except ImportError as exc:  # pragma: no cover
    sys.exit(
        f"Não consegui importar drive_service.py ({exc}). "
        "Rode este script na pasta do projeto."
    )

DATABASE_PADRAO = os.path.join(BASE_DIR, "database", "database", "alunosv3.db")
PLANILHA_PADRAO = os.path.join(BASE_DIR, "sincronizacao_drive.xlsx")

CAMPOS_UPLOAD = [
    "doc_cpf",
    "doc_cartao_sus",
    "doc_certidao",
    "doc_cpf_mae",
    "doc_cpf_pai",
    "doc_cep",
    "doc_laudo",
    "doc_carteira_vacinacao",
    "doc_comprovante_vacinacao",
]
PASTA_MIME = "application/vnd.google-apps.folder"


# ── Normalização ──────────────────────────────────────────────────────────────
def norm(s) -> str:
    """MAIÚSCULAS, sem acentos, espaços colapsados — para comparar nomes."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


# subpasta (normalizada) → campos que ela pode conter
PASTA_PARA_CAMPOS: dict[str, list[str]] = {}
for _campo, _pasta in drive.CAMPO_PARA_SUBPASTA.items():
    if _campo in CAMPOS_UPLOAD:
        PASTA_PARA_CAMPOS.setdefault(norm(_pasta), []).append(_campo)

_MAE = {"mae", "materno", "materna", "genitora"}
_PAI = {"pai", "paterno", "genitor"}


def _vac(t: set) -> bool:
    return any(x.startswith("vacin") for x in t)


# Regras sobre as palavras do NOME DO ARQUIVO (usadas quando a subpasta guarda
# mais de um tipo de documento, ex.: "RG E CERT_NASCIMENTO").
REGRAS = {
    "doc_cep": lambda t: bool(
        t & {"comprovante", "residencia", "endereco", "luz", "agua", "energia", "cep"}
    ),
    "doc_certidao": lambda t: any(x.startswith("cert") for x in t) or "nascimento" in t,
    "doc_cpf": lambda t: bool(t & {"rg", "identidade", "cpf", "cnh"})
    and not (t & (_MAE | _PAI)),
    "doc_cpf_mae": lambda t: bool(t & _MAE),
    "doc_cpf_pai": lambda t: bool(t & _PAI),
    "doc_laudo": lambda t: bool(t & {"laudo", "medico", "cid"}),
    "doc_cartao_sus": lambda t: "sus" in t and not _vac(t),
    "doc_comprovante_vacinacao": lambda t: _vac(t) and "comprovante" in t,
    "doc_carteira_vacinacao": lambda t: (_vac(t) and "comprovante" not in t)
    or "caderneta" in t
    or ("carteira" in t and "sus" not in t),
}


def inferir_campo(subpasta: str, arquivo: str):
    """Devolve (campo, motivo_se_nao_deu). Só decide quando há UMA resposta clara."""
    candidatos = PASTA_PARA_CAMPOS.get(norm(subpasta)) if subpasta else None
    if candidatos and len(candidatos) == 1:
        return candidatos[0], ""
    universo = candidatos or CAMPOS_UPLOAD
    base = os.path.splitext(arquivo)[0]
    tokens = set(re.sub(r"[^a-z0-9]+", " ", norm(base).lower()).split())
    achados = [c for c in universo if REGRAS[c](tokens)]
    if len(achados) == 1:
        return achados[0], ""
    if not achados:
        return None, "nome do arquivo não indica o tipo do documento"
    return None, "nome do arquivo combina com mais de um tipo"


# ── Varredura do Drive ────────────────────────────────────────────────────────
_CAMPOS_API = "nextPageToken, files(id, name, mimeType, parents, webViewLink)"


def listar_filhos(service, pais: list) -> dict:
    """Filhos (não-lixeira) de várias pastas de uma vez: {id_pai: [itens]}."""
    filhos = {p: [] for p in pais}
    for i in range(0, len(pais), 20):  # 20 pastas por consulta = poucas chamadas
        lote = pais[i : i + 20]
        q = "(" + " or ".join(f"'{p}' in parents" for p in lote) + ") and trashed=false"
        token = None
        while True:
            res = (
                service.files()
                .list(
                    q=q,
                    spaces="drive",
                    fields=_CAMPOS_API,
                    pageSize=1000,
                    pageToken=token,
                    supportsAllDrives=True,
                    includeItemsFromAllDrives=True,
                )
                .execute()
            )
            for f in res.get("files", []):
                for p in f.get("parents", []):
                    if p in filhos:
                        filhos[p].append(f)
            token = res.get("nextPageToken")
            if not token:
                break
    return filhos


def varrer_drive(service, root_id: str):
    """Percorre ANO/SALA/ALUNOS/NOME/SUBPASTA e devolve (registros, nº pastas aninhadas ignoradas)."""

    def so_pastas(itens):
        return [i for i in itens if i["mimeType"] == PASTA_MIME]

    print("  lendo anos...")
    anos = so_pastas(listar_filhos(service, [root_id])[root_id])
    f_anos = listar_filhos(service, [a["id"] for a in anos])
    salas = [(a["name"], s) for a in anos for s in so_pastas(f_anos[a["id"]])]

    print(f"  lendo {len(salas)} sala(s)...")
    f_salas = listar_filhos(service, [s["id"] for _, s in salas])
    conts = [
        (ano, s["name"], c)
        for ano, s in salas
        for c in so_pastas(f_salas[s["id"]])
        if norm(c["name"]) == "ALUNOS"
    ]

    f_conts = listar_filhos(service, [c["id"] for _, _, c in conts])
    pastas = [
        (ano, sala, p) for ano, sala, c in conts for p in so_pastas(f_conts[c["id"]])
    ]
    print(f"  lendo {len(pastas)} pasta(s) de aluno...")

    f_pastas = listar_filhos(service, [p["id"] for _, _, p in pastas])
    registros, subs, aninhadas = [], [], 0

    def novo(ano, sala, pasta, subpasta, arq):
        return {
            "ano": ano,
            "sala": sala,
            "pasta_aluno": pasta["name"],
            "subpasta": subpasta,
            "arquivo": arq["name"],
            "drive_id": arq["id"],
            "drive_url": arq.get("webViewLink", ""),
        }

    for ano, sala, p in pastas:
        for it in f_pastas[p["id"]]:
            if it["mimeType"] == PASTA_MIME:
                subs.append((ano, sala, p, it))
            else:
                registros.append(novo(ano, sala, p, "", it))

    f_subs = listar_filhos(service, [s["id"] for *_, s in subs])
    for ano, sala, p, sub in subs:
        for it in f_subs[sub["id"]]:
            if it["mimeType"] == PASTA_MIME:
                aninhadas += 1
            else:
                registros.append(novo(ano, sala, p, sub["name"], it))
    return registros, aninhadas


# ── Casamento pasta ↔ aluno ───────────────────────────────────────────────────
def indexar_alunos(con) -> dict:
    idx: dict[str, list] = {}
    for r in con.execute(
        "SELECT a.id, a.nome_completo, m.ano, m.sala FROM aluno a "
        "LEFT JOIN matricula m ON m.aluno_id = a.id"
    ):
        idx.setdefault(norm(r["nome_completo"]), []).append(dict(r))
    return idx


def achar_aluno(idx: dict, pasta_aluno: str, ano: str, sala: str):
    """(aluno|None, status, observação)."""
    chave = norm(pasta_aluno)
    cands = idx.get(chave, [])
    if len(cands) > 1:  # homônimos/duplicatas: desempata por ano e sala da pasta
        melhores = [
            c
            for c in cands
            if norm(c["ano"]) == norm(ano) and norm(c["sala"]) == norm(sala)
        ]
        if len(melhores) == 1:
            return melhores[0], "", ""
        ids = ", ".join(str(c["id"]) for c in cands)
        return (
            None,
            "ALUNO AMBÍGUO",
            f"há {len(cands)} alunos com este nome (ids {ids}); informe o aluno_id",
        )
    if len(cands) == 1:
        return cands[0], "", ""
    parecido = difflib.get_close_matches(chave, list(idx), n=1, cutoff=0.85)
    obs = "nenhum aluno com este nome no sistema"
    if parecido:
        c = idx[parecido[0]][0]
        obs += f" — parecido: {c['nome_completo']} (id {c['id']})"
    return None, "SEM ALUNO", obs


# ── Planilha de conferência ───────────────────────────────────────────────────
COLUNAS = [
    "status",
    "aluno_id",
    "aluno_no_sistema",
    "pasta_aluno",
    "ano",
    "sala",
    "subpasta",
    "arquivo",
    "campo",
    "observacao",
    "drive_id",
    "drive_url",
]
EDITAVEIS = {"aluno_id", "campo"}
ORDEM_STATUS = {"SEM ALUNO": 0, "ALUNO AMBÍGUO": 1, "REVISAR": 2, "OK": 3}


def gerar_planilha(linhas: list, caminho: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos no Drive"
    ws.append(COLUNAS)
    cores = {
        "OK": "C6EFCE",
        "REVISAR": "FFEB9C",
        "SEM ALUNO": "FFC7CE",
        "ALUNO AMBÍGUO": "F8CBAD",
    }
    for c, nome in enumerate(COLUNAS, 1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(bold=True)
        cel.fill = PatternFill(
            "solid", fgColor="FFFF00" if nome in EDITAVEIS else "D9D9D9"
        )

    for l in sorted(
        linhas,
        key=lambda x: (ORDEM_STATUS[x["status"]], norm(x["pasta_aluno"]), x["arquivo"]),
    ):
        ws.append([l.get(c, "") for c in COLUNAS])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill(
            "solid", fgColor=cores[l["status"]]
        )

    larg = {
        "status": 15,
        "aluno_id": 9,
        "aluno_no_sistema": 32,
        "pasta_aluno": 32,
        "ano": 7,
        "sala": 10,
        "subpasta": 34,
        "arquivo": 38,
        "campo": 26,
        "observacao": 55,
        "drive_id": 14,
        "drive_url": 14,
    }
    for i, nome in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg[nome]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(ws.max_row, 2)}"

    dv = DataValidation(
        type="list", formula1='"' + ",".join(CAMPOS_UPLOAD) + '"', allow_blank=True
    )
    ws.add_data_validation(dv)
    dv.add(
        f"{get_column_letter(COLUNAS.index('campo') + 1)}2:{get_column_letter(COLUNAS.index('campo') + 1)}{max(ws.max_row, 2)}"
    )

    lg = wb.create_sheet("Legenda")
    lg.append(["Como usar"])
    for t in [
        "1) Colunas amarelas (aluno_id e campo) são as que você pode preencher/corrigir.",
        "2) OK = aluno e tipo de documento identificados: será gravado como está.",
        "3) REVISAR = a pasta tem mais de um tipo de documento e o nome do arquivo não deixa claro: escolha o 'campo' na lista.",
        "4) SEM ALUNO / ALUNO AMBÍGUO = digite o aluno_id correto (o id aparece na coluna 'observacao' quando há um parecido).",
        "5) Linha sem aluno_id ou sem campo NÃO é gravada. Para ignorar um arquivo, deixe o campo em branco.",
        "6) Salve a planilha e rode:  python sincronizar_drive.py --aplicar",
    ]:
        lg.append([t])
    lg.append([])
    lg.append(["campo", "o que é", "subpasta no Drive"])
    for c in CAMPOS_UPLOAD:
        lg.append(
            [
                c,
                drive.CAMPO_NOME_DISPLAY.get(c, c),
                drive.CAMPO_PARA_SUBPASTA.get(c, ""),
            ]
        )
    lg.column_dimensions["A"].width = 28
    lg.column_dimensions["B"].width = 30
    lg.column_dimensions["C"].width = 45
    wb.save(caminho)


# ── Banco ─────────────────────────────────────────────────────────────────────
def abrir_banco(caminho: str):
    if not os.path.exists(caminho):
        sys.exit(
            f"Banco não encontrado: {caminho}\nUse --banco para informar o caminho."
        )
    con = sqlite3.connect(caminho, timeout=30, isolation_level=None)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA busy_timeout = 30000")
    con.execute("PRAGMA foreign_keys = ON")
    cols = {r["name"] for r in con.execute("PRAGMA table_info(documento)")}
    faltam = {"aluno_id", "campo", "nome", "subpasta", "drive_id", "drive_url"} - cols
    if faltam:
        sys.exit(
            f"A tabela 'documento' não tem as colunas {sorted(faltam)}. "
            "Abra o sistema uma vez (python app.py) para atualizar o banco e tente de novo."
        )
    return con


# ── Comandos ──────────────────────────────────────────────────────────────────
def simular(args):
    root_id = os.getenv("DRIVE_ROOT_FOLDER_ID")
    if not root_id:
        sys.exit("DRIVE_ROOT_FOLDER_ID não está no .env.")
    try:
        service = drive.get_drive_service()
    except Exception as exc:
        sys.exit(
            f"Drive não autenticado: {exc}\nRode 'python autenticar_drive.py' e tente de novo."
        )

    con = abrir_banco(args.banco)
    print("Lendo o Google Drive...")
    registros, aninhadas = varrer_drive(service, root_id)

    ja = {
        r[0]
        for r in con.execute(
            "SELECT drive_id FROM documento WHERE drive_id IS NOT NULL"
        )
    }
    idx = indexar_alunos(con)

    linhas, ja_cadastrados, com_arquivo = [], 0, set()
    for r in registros:
        if r["drive_id"] in ja:
            ja_cadastrados += 1
            continue
        aluno, status_aluno, obs = achar_aluno(
            idx, r["pasta_aluno"], r["ano"], r["sala"]
        )
        campo, motivo = inferir_campo(r["subpasta"], r["arquivo"])
        if not r["subpasta"]:
            obs = (obs + "; " if obs else "") + "arquivo solto, fora de subpasta"
        if not campo and motivo:
            obs = (obs + "; " if obs else "") + motivo
        if aluno:
            com_arquivo.add(aluno["id"])
        status = status_aluno or ("OK" if campo else "REVISAR")
        linhas.append(
            {
                **r,
                "status": status,
                "aluno_id": aluno["id"] if aluno else "",
                "aluno_no_sistema": aluno["nome_completo"] if aluno else "",
                "campo": campo or "",
                "observacao": obs,
            }
        )

    gerar_planilha(linhas, args.planilha)

    contagem = {s: sum(1 for l in linhas if l["status"] == s) for s in ORDEM_STATUS}
    total_alunos = con.execute("SELECT COUNT(*) FROM aluno").fetchone()[0]
    print("\n=== RESULTADO DA SIMULAÇÃO (nada foi gravado no banco) ===")
    print(f"Arquivos encontrados no Drive ............ {len(registros)}")
    print(f"  já registrados no sistema (ignorados) .. {ja_cadastrados}")
    print(f"  a registrar ............................ {len(linhas)}")
    for s in ORDEM_STATUS:
        print(f"     {s:<15} {contagem[s]}")
    print(
        f"Alunos do sistema com arquivos no Drive .. {len(com_arquivo)} de {total_alunos}"
    )
    if aninhadas:
        print(f"Atenção: {aninhadas} pasta(s) dentro de subpastas foram ignoradas.")
    print(f"\nPlanilha gerada: {args.planilha}")
    pend = contagem["SEM ALUNO"] + contagem["ALUNO AMBÍGUO"] + contagem["REVISAR"]
    if pend:
        print(
            f"Abra-a, resolva as {pend} linha(s) amarelas/vermelhas e salve; depois rode com --aplicar."
        )
    else:
        print(
            "Tudo identificado. Se estiver de acordo, rode:  python sincronizar_drive.py --aplicar"
        )


def aplicar(args):
    from openpyxl import load_workbook

    if not os.path.exists(args.planilha):
        sys.exit(
            f"Planilha não encontrada: {args.planilha}\nRode primeiro sem --aplicar."
        )
    wb = load_workbook(args.planilha, data_only=True)
    ws = (
        wb["Documentos no Drive"]
        if "Documentos no Drive" in wb.sheetnames
        else wb.worksheets[0]
    )
    cab = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
    for obrig in ("aluno_id", "campo", "drive_id", "arquivo", "subpasta"):
        if obrig not in cab:
            sys.exit(
                f"Coluna '{obrig}' não encontrada na planilha. Gere-a de novo (sem --aplicar)."
            )

    def val(row, nome):
        v = row[cab[nome]].value if nome in cab and cab[nome] < len(row) else None
        return "" if v is None else str(v).strip()

    con = abrir_banco(args.banco)
    ids_validos = {r[0] for r in con.execute("SELECT id FROM aluno")}
    ja = {
        r[0]
        for r in con.execute(
            "SELECT drive_id FROM documento WHERE drive_id IS NOT NULL"
        )
    }

    novos, sem_dados, invalidos, repetidos = [], 0, 0, 0
    vistos = set()
    for row in ws.iter_rows(min_row=2):
        drive_id, campo, aid = (
            val(row, "drive_id"),
            val(row, "campo"),
            val(row, "aluno_id"),
        )
        if not drive_id:
            continue
        if not aid or not campo:
            sem_dados += 1
            continue
        try:
            aid = int(float(aid))
        except ValueError:
            invalidos += 1
            continue
        if campo not in CAMPOS_UPLOAD or aid not in ids_validos:
            invalidos += 1
            continue
        if drive_id in ja or drive_id in vistos:
            repetidos += 1
            continue
        vistos.add(drive_id)
        novos.append(
            (
                aid,
                campo,
                val(row, "arquivo"),
                drive.CAMPO_PARA_SUBPASTA.get(campo, val(row, "subpasta")),
                drive_id,
                val(row, "drive_url") or None,
            )
        )

    if not novos:
        print("Nada novo para gravar.")
    else:
        pasta_bkp = os.path.join(os.path.dirname(args.banco), "backups")
        os.makedirs(pasta_bkp, exist_ok=True)
        bkp = os.path.join(
            pasta_bkp, f"antes_sync_drive_{datetime.now():%Y%m%d_%H%M%S}.db"
        )
        dest = sqlite3.connect(bkp)
        con.backup(dest)
        dest.close()
        print(f"Backup do banco: {bkp}")

        con.execute("BEGIN IMMEDIATE")
        try:
            con.executemany(
                "INSERT INTO documento (aluno_id, campo, nome, subpasta, drive_id, drive_url) "
                "VALUES (?,?,?,?,?,?)",
                novos,
            )
            con.execute("COMMIT")
        except BaseException:
            if con.in_transaction:
                con.execute("ROLLBACK")
            raise

    print("\n=== RESULTADO ===")
    print(f"Documentos gravados ........................ {len(novos)}")
    print(f"Ignorados (já estavam registrados) ......... {repetidos}")
    print(f"Ignorados (sem aluno_id ou sem campo) ...... {sem_dados}")
    print(f"Ignorados (aluno_id/campo inválido) ........ {invalidos}")
    if novos:
        alunos_ok = len({n[0] for n in novos})
        print(
            f"\nAlunos que ganharam documentos: {alunos_ok}. Atualize a tela de Documentação no sistema."
        )


def main():
    ap = argparse.ArgumentParser(
        description="Registra no sistema os documentos que já estão no Drive."
    )
    ap.add_argument(
        "--aplicar",
        action="store_true",
        help="grava no banco o conteúdo da planilha revisada",
    )
    ap.add_argument("--planilha", default=PLANILHA_PADRAO)
    ap.add_argument("--banco", default=DATABASE_PADRAO)
    args = ap.parse_args()
    aplicar(args) if args.aplicar else simular(args)


if __name__ == "__main__":
    main()
